"""Read a bounded .xlsx upload in memory; do not evaluate formulas or make requests."""
from __future__ import annotations

from datetime import date, datetime, time
from io import BytesIO
import math
from zipfile import BadZipFile, ZipFile

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 25 * 1024 * 1024
MAX_ZIP_MEMBERS = 1000
MAX_ROWS = 5000
MAX_COLUMNS = 20


def read_workbook(body):
    """Return the first sheet, including headers and row positions, as JSON cells."""
    if not body or len(body) > MAX_UPLOAD_BYTES:
        raise ValueError("Choose a non-empty .xlsx file no larger than 5 MB.")
    try:
        with ZipFile(BytesIO(body)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ZIP_MEMBERS or sum(item.file_size for item in members) > MAX_UNCOMPRESSED_BYTES:
                raise ValueError("The workbook is too large to preview. Export only the concentration table.")
            names = {item.filename.lower() for item in members}
            if "[content_types].xml" not in names or "xl/workbook.xml" not in names:
                raise ValueError("Choose an Excel .xlsx workbook.")
            if any(name.endswith("vbaproject.bin") for name in names):
                raise ValueError("Save this workbook as .xlsx without macros before importing.")
            if any(item.flag_bits & 1 for item in members):
                raise ValueError("Remove the workbook password before importing.")
    except BadZipFile:
        raise ValueError("This file could not be read as .xlsx. Save it as .xlsx and try again.") from None

    from openpyxl import load_workbook
    workbook = None
    try:
        workbook = load_workbook(BytesIO(body), read_only=True, data_only=True, keep_links=False)
        if not workbook.worksheets:
            raise ValueError("The workbook does not contain a worksheet.")
        sheet = workbook.worksheets[0]
        if (sheet.max_row or 0) > MAX_ROWS or (sheet.max_column or 0) > MAX_COLUMNS:
            raise ValueError("The first worksheet must contain at most 5,000 rows and 20 columns.")
        rows = []
        # Stream actual rows rather than trusting a possibly inaccurate dimension tag.
        sheet.reset_dimensions()
        for raw in sheet.iter_rows(values_only=True):
            if len(rows) == MAX_ROWS:
                raise ValueError("The first worksheet must contain at most 5,000 rows.")
            if len(raw) > MAX_COLUMNS:
                raise ValueError("The first worksheet must contain at most 20 columns.")
            cells = []
            for value in raw[:MAX_COLUMNS]:
                if isinstance(value, (datetime, date, time)):
                    value = value.isoformat()
                elif isinstance(value, float) and not math.isfinite(value):
                    value = str(value)
                elif value is not None and not isinstance(value, (str, int, float, bool)):
                    value = str(value)
                cells.append(value)
            while cells and cells[-1] is None:
                cells.pop()
            rows.append(cells)
        while rows and not rows[-1]:
            rows.pop()
        if not rows:
            raise ValueError("The first worksheet is empty.")
        return {"sheet": sheet.title, "rows": rows}
    except ValueError:
        raise
    except Exception:
        raise ValueError("The workbook could not be read. Save it as a new .xlsx file and try again.") from None
    finally:
        if workbook is not None:
            workbook.close()
