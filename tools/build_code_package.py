"""Repackage the relative local_code artifacts without changing user data.

Use --output-dir for a staging destination. Both output files are fully prepared
before publication; each replacement is atomic and the manifest is published
last. A caught publication error restores the previous pair.
"""
from __future__ import annotations
import argparse
import csv
import hashlib
import json
import os
import shutil
import tempfile
from datetime import datetime, timezone
from pathlib import Path
import zipfile

ROOT = Path(__file__).resolve().parents[1]
PACKAGE = ROOT / 'local_code'
VERSION = '2026-10-10-cpu-2'
ARCHIVE_NAME = 'WPN-local-code.zip'
MANIFEST_NAME = 'WPN-local-code.manifest.json'
ROOT_FILES = (
    'README.md', 'NOTICE.txt', 'requirements.txt', 'requirements-tested.txt',
    'Setup_WPN.bat', 'Run_WPN.bat', 'run.py', 'wpn_core.py',
    'code_1_run_wpn.py', 'code_2_plot_results.py',
)
MODEL_FILES = (
    'model.joblib', 'imputer.joblib', 'scaler.joblib', 'feature_columns.csv',
    'feature_preprocessing.csv', 'tabpfn-v2.6-regressor-v2.6_default.ckpt',
)
SUPPORT_FILES = (
    'data/chemicals.json', 'data/parameters.csv', 'data/chemical_list.csv',
    'examples/example_52.csv', 'examples/example_52.xlsx',
    'licenses/TabPFN-python-LICENSE.txt', 'licenses/TabPFN-2.6-model-LICENSE.txt',
    'results/README.txt',
)

def sha256(path):
    result = hashlib.sha256()
    with Path(path).open('rb') as handle:
        for data in iter(lambda: handle.read(1024 * 1024), b''):
            result.update(data)
    return result.hexdigest()

def read_csv(path):
    with Path(path).open(encoding='utf-8-sig', newline='') as handle:
        return list(csv.DictReader(handle))

def table(path, rows):
    with path.open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.writer(stream)
        writer.writerow(['Chemical', 'WW (nM)'])
        writer.writerows(rows)

def workbook(path, rows):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation
    book = Workbook()
    sheet = book.active
    sheet.title = 'Concentrations'
    sheet.append(['Chemical', 'WW (nM)'])
    for row in rows:
        sheet.append(list(row))
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:B{sheet.max_row}'
    sheet.column_dimensions['A'].width = 64
    sheet.column_dimensions['B'].width = 22
    for cell in sheet[1]:
        cell.font = Font(color='FFFFFF', bold=True)
        cell.fill = PatternFill('solid', fgColor='255BE3')
    validation = DataValidation(type='decimal', operator='greaterThan', formula1='0', allow_blank=True)
    validation.promptTitle = 'WW concentration (nM)'
    validation.prompt = 'Enter a positive number. Leave unmeasured chemicals blank.'
    validation.error = 'Enter a positive concentration in nM.'
    validation.showErrorMessage = True
    validation.showInputMessage = True
    sheet.add_data_validation(validation)
    validation.add('B2:B5000')
    guide = book.create_sheet('Read me')
    for text in ['Fill WW (nM) in the first sheet.','Use chemical name, CAS or InChIKey from the supplied 52-chemical catalog.','Fill only the compounds you measured; leave the others blank or delete their rows.','Save and close this file, then run Run_WPN.bat.','Only valid, unique chemicals are calculated. See results/*/input_check.csv.']:
        guide.append([text])
    guide.column_dimensions['A'].width = 105
    for row in guide:
        row[0].alignment = Alignment(wrap_text=True)
    book.save(path)

def source_file(package, relative):
    path = package / relative
    if not path.resolve().is_relative_to(package):
        raise ValueError(f'Package file points outside local_code: {relative}')
    if not path.is_file():
        raise FileNotFoundError(f'Required package file is missing: {relative}')
    return path


def package_sources(package):
    """Validate local artifacts and return only the files needed by the runner."""
    manifest_path = source_file(package, 'models/manifest.json')
    manifest = json.loads(manifest_path.read_text(encoding='utf-8-sig'))
    hashes = manifest['sha256']
    if not isinstance(hashes, dict) or not set(MODEL_FILES).issubset(hashes):
        raise ValueError('Model manifest is missing required artifact hashes.')
    files = [*ROOT_FILES, *SUPPORT_FILES, 'models/manifest.json']
    for name, expected in hashes.items():
        if not name or '/' in name or '\\' in name or ':' in name or name in {'.', '..'}:
            raise ValueError(f'Unsafe model artifact name: {name}')
        relative = f'models/{name}'
        if sha256(source_file(package, relative)) != expected:
            raise ValueError(f'Model artifact hash does not match: {name}')
        files.append(relative)
    sources = {relative: source_file(package, relative) for relative in files}
    catalog = json.loads(sources['data/chemicals.json'].read_text(encoding='utf-8-sig'))
    identifiers = {row['id'] for row in catalog}
    if len(catalog) != 52 or len(identifiers) != 52:
        raise ValueError('The local chemical catalog must contain 52 unique chemicals.')
    for relative, key in [('data/parameters.csv', 'StdInChIKey'), ('data/chemical_list.csv', 'InChIKey')]:
        rows = read_csv(sources[relative])
        if len(rows) != 52 or {row[key] for row in rows} != identifiers:
            raise ValueError(f'{relative} does not match the 52-chemical catalog.')
    return sources, sorted(catalog, key=lambda row: row['name'].lower()), manifest


def publish_pair(staged, destination):
    """Replace complete files, with the manifest last and rollback on failure.

    Two separate paths cannot be replaced as one filesystem transaction. The
    manifest checksum lets consumers detect the brief interval between replaces.
    """
    names = (ARCHIVE_NAME, MANIFEST_NAME)
    previous = {}
    for name in names:
        target = destination / name
        if target.exists():
            backup = staged / (name + '.previous')
            shutil.copy2(target, backup)
            previous[name] = backup
    replaced = []
    try:
        for name in names:
            os.replace(staged / name, destination / name)
            replaced.append(name)
    except OSError:
        for name in reversed(replaced):
            target = destination / name
            if name in previous:
                os.replace(previous[name], target)
            else:
                target.unlink()
        raise


def build_package(package=PACKAGE, output_dir=ROOT / 'public/downloads'):
    package = Path(package).resolve()
    output_dir = Path(output_dir).resolve()
    if output_dir.is_relative_to(package):
        raise ValueError('Output directory must be outside local_code to preserve user files.')
    sources, catalog, model_manifest = package_sources(package)
    output_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.wpn-package-', dir=output_dir) as folder:
        staged = Path(folder)
        inputs = staged / 'inputs'
        inputs.mkdir()
        blank = [(row['name'], None) for row in catalog]
        workbook(inputs / 'concentrations.xlsx', blank)
        table(inputs / 'concentrations.csv', blank)
        sources.update({f'inputs/{path.name}': path for path in inputs.iterdir()})
        output = staged / ARCHIVE_NAME
        with zipfile.ZipFile(output, 'w', compression=zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
            for relative, path in sorted(sources.items()):
                archive.write(path, f'WPN-local-code/{relative}')
        info = {
            'filename': ARCHIVE_NAME, 'bytes': output.stat().st_size,
            'sha256': sha256(output), 'version': VERSION,
            'modelVersion': model_manifest['version'],
            'builtAt': datetime.now(timezone.utc).isoformat(),
            'inputFile': 'inputs/concentrations.xlsx', 'runFile': 'Run_WPN.bat',
            'setupFile': 'Setup_WPN.bat', 'python': '3.12', 'device': 'CPU',
        }
        (staged / MANIFEST_NAME).write_text(json.dumps(info, indent=2), encoding='utf-8')
        publish_pair(staged, output_dir)
    return info


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--package-dir', type=Path, default=PACKAGE,
                        help='Existing local_code directory (read only).')
    parser.add_argument('--output-dir', type=Path, default=ROOT / 'public/downloads',
                        help='Destination for ZIP and manifest; use a staging directory to preview.')
    args = parser.parse_args()
    print(json.dumps(build_package(args.package_dir, args.output_dir), indent=2))

if __name__ == '__main__':
    main()
